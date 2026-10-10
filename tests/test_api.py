"""API and ingestion behaviour on both backends: upload safety, idempotent re-index, scanned PDFs,
metrics, incremental ingestion, and (pgvector) several API replicas sharing one database."""
import io
import os

import pytest
from fastapi.testclient import TestClient

import ingest
import pipeline
from app import create_app
from conftest import PG_DSN, fake_embed
from store import VectorLiteStore


@pytest.fixture
def env(make_store, tmp_path):
    docs = tmp_path / "docs"
    store = make_store()
    app = create_app(embedder=fake_embed, docs_dir=str(docs), store=store)
    with TestClient(app) as client:
        yield client, tmp_path, docs, store


def upload(client, name, content: bytes):
    return client.post("/upload", files={"file": (name, io.BytesIO(content), "application/octet-stream")})


def blank_pdf(pages=2) -> bytes:
    from pypdf import PdfWriter
    w = PdfWriter()
    for _ in range(pages):
        w.add_blank_page(width=612, height=792)
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


# --- upload safety ---

def test_path_traversal_is_contained(env):
    client, root, docs, store = env
    r = upload(client, "../../escape.txt", b"graduate student researcher stipend")
    assert r.status_code == 200, r.text
    assert not (root / "escape.txt").exists() and not (root.parent / "escape.txt").exists()
    assert "escape.txt" in store.file_stats()
    if store.stores_sources:
        assert store.source("escape.txt") == b"graduate student researcher stipend"
        assert not (docs / "escape.txt").exists(), "pgvector mode must not depend on local disk"
    else:
        assert (docs / "escape.txt").exists()


@pytest.mark.parametrize("name", ["notes.exe", "....", "../", ".pdf"])
def test_rejects_bad_names(env, name):
    client, _, docs, store = env
    assert upload(client, name, b"text").status_code == 400
    assert list(docs.iterdir()) == [] and len(store) == 0


def test_rejects_oversized_upload(env, monkeypatch):
    client, *_ = env
    monkeypatch.setattr(pipeline, "MAX_UPLOAD_BYTES", 10)
    assert upload(client, "big.txt", b"x" * 11).status_code == 413


# --- re-upload behaviour ---

def test_reupload_identical_is_noop(env):
    client, *_ = env
    body = b"fee remission " * 200  # several chunks
    first = upload(client, "policy.txt", body).json()
    second = upload(client, "policy.txt", body).json()
    assert second["unchanged"] is True
    assert second["chunks"] == first["chunks"]
    assert client.get("/health").json()["vectors"] == first["chunks"]


def test_reupload_changed_replaces_old_chunks(env):
    client, *_ = env
    long = upload(client, "policy.txt", b"tuition " * 500).json()
    short = upload(client, "policy.txt", b"nonresident supplemental tuition waived").json()
    assert long["chunks"] > 1 and short["chunks"] == 1
    assert client.get("/health").json()["vectors"] == 1
    files = {f["filename"]: f for f in client.get("/files").json()}
    assert files["policy.txt"]["count"] == 1


# --- extraction ---

def test_scanned_pdf_reports_reason(env):
    client, _, docs, store = env
    r = upload(client, "scan.pdf", blank_pdf(3))
    assert r.status_code == 422
    assert "3 of 3 pages have no text layer" in r.json()["error"]
    assert not (docs / "scan.pdf").exists() and "scan.pdf" not in store.file_stats()


def test_corrupt_pdf_is_client_error(env):
    client, *_ = env
    assert upload(client, "broken.pdf", b"%PDF-not really").status_code == 422


def test_docx_tables_are_extracted():
    from docx import Document
    d = Document()
    d.add_paragraph("Stipend schedule")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "GSR Step 6", "$4,000"
    buf = io.BytesIO()
    d.save(buf)
    assert "GSR Step 6 | $4,000" in pipeline.extract_text(buf.getvalue(), "s.docx").text


def test_sidecar_txt_is_not_indexed_twice(tmp_path):
    (tmp_path / "a.pdf").write_bytes(b"%PDF")
    (tmp_path / "a.pdf.txt").write_text("copy")
    (tmp_path / "notes.txt").write_text("real")
    names = [os.path.basename(p) for p in pipeline.list_source_files(str(tmp_path))]
    assert names == ["a.pdf", "notes.txt"]


# --- search, metrics, files ---

def test_search_metrics_split_stages(env):
    client, *_ = env
    upload(client, "a.txt", b"nonresident supplemental tuition")
    upload(client, "b.txt", b"teaching assistant appointment hours")
    hits = client.get("/search", params={"q": "nonresident tuition", "k": 1}).json()
    assert hits[0]["metadata"]["file"] == "a.txt"
    filtered = client.get("/search", params={"q": "nonresident tuition", "k": 5, "file": "b.txt"}).json()
    assert {h["metadata"]["file"] for h in filtered} == {"b.txt"}
    m = client.get("/metrics").json()
    assert m["count"] == 2
    for key in ("p50_ms", "embed_p50_ms", "search_p50_ms"):
        assert m[key] is not None
    assert m["last_query"]["latency_ms"] == pytest.approx(
        m["last_query"]["embed_ms"] + m["last_query"]["search_ms"], abs=0.01)


def test_health_reports_backend(env):
    client, _, _, store = env
    assert client.get("/health").json()["store"] == store.name


def test_bench_and_parity(env):
    client, *_ = env
    b = client.get("/bench", params={"N": 300, "k": 5}).json()
    assert b["N"] == 300 and b["file_MB"] > 0
    assert client.get("/parity", params={"K": 5}).json()["ok"] is True


# --- several API replicas on one database (pgvector only) ---

def test_replicas_share_one_database(pg_factory, tmp_path):
    """Two independent app instances, as two Kubernetes pods would be, each with its own local disk."""
    a = create_app(embedder=fake_embed, docs_dir=str(tmp_path / "pod-a"), store=pg_factory())
    b = create_app(embedder=fake_embed, docs_dir=str(tmp_path / "pod-b"), store=pg_factory())
    with TestClient(a) as pod_a, TestClient(b) as pod_b:
        assert upload(pod_a, "gsr.txt", b"graduate student researcher appointment").status_code == 200
        hits = pod_b.get("/search", params={"q": "graduate student researcher", "k": 1}).json()
        assert hits[0]["metadata"]["file"] == "gsr.txt"
        assert pod_b.get("/files").json() == pod_a.get("/files").json()
        again = upload(pod_b, "gsr.txt", b"graduate student researcher appointment").json()
        assert again["unchanged"] is True, "pod B must see pod A's upload as the same content"


# --- incremental ingestion ---

def test_ingest_vectorlite_is_incremental_and_prunes(tmp_path, capsys):
    docs, db = tmp_path / "docs", tmp_path / "kb.db"
    docs.mkdir()
    (docs / "a.txt").write_text("graduate funding " * 100)
    (docs / "b.txt").write_text("fellowship rules")
    args = ["--docs", str(docs), "--db", str(db), "--store", "vectorlite"]

    ingest.main(args, embedder=fake_embed)
    assert "2 indexed" in capsys.readouterr().out

    calls = []
    ingest.main(args, embedder=lambda t: calls.append(t) or fake_embed(t))
    assert calls == [] and "2 unchanged" in capsys.readouterr().out

    (docs / "b.txt").unlink()
    (docs / "a.txt").write_text("updated policy")
    ingest.main(args, embedder=fake_embed)
    out = capsys.readouterr().out
    assert "1 indexed" in out and "1 removed" in out
    assert set(VectorLiteStore(str(db), dimension=pipeline.EMBED_DIM).file_stats()) == {"a.txt"}


def test_ingest_pgvector_keeps_uploads_unless_pruning(pg_factory, tmp_path, monkeypatch, capsys):
    monkeypatch.setenv("VB_PG_DSN", PG_DSN)
    monkeypatch.setenv("VB_PG_SCHEMA", pg_factory.schema)
    store = pg_factory()
    store.replace_file("uploaded.txt", ["only in the database"], fake_embed(["only in the database"]),
                       sha256="u", source=b"only in the database")
    docs = tmp_path / "docs"
    docs.mkdir()
    (docs / "a.txt").write_text("graduate funding")
    args = ["--docs", str(docs), "--store", "pgvector"]

    ingest.main(args, embedder=fake_embed)
    out = capsys.readouterr().out
    assert "1 indexed" in out and "0 removed" in out
    assert set(store.file_stats()) == {"a.txt", "uploaded.txt"}
    assert store.source("a.txt") == b"graduate funding"

    ingest.main(args + ["--prune"], embedder=fake_embed)
    assert "1 removed" in capsys.readouterr().out
    assert set(store.file_stats()) == {"a.txt"}

    with pytest.raises(SystemExit):
        ingest.main(args + ["--rebuild"], embedder=fake_embed)
